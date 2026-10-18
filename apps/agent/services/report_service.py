"""Real ledger reports. A send is Gmail acceptance, never proof of delivery/opening."""

from collections import Counter
from datetime import date, datetime, time, timedelta, timezone
from hashlib import sha256
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo

from core.store import store

TZ = ZoneInfo("Africa/Johannesburg")
OUTCOMES = ("unclassified", "interested", "declined", "won", "unsubscribed")
DEFINITIONS = {
    "scope": "Only messages recorded by this installation; connection tests are excluded unless explicitly included.",
    "sent": "Gmail-confirmed submissions, dated by the recorded submission attempt. Acceptance does not prove delivery or opening.",
    "cohort": "Conversations first sent in the selected dates. One conversation per mailbox/thread; repeated messages do not inflate the denominator.",
    "reply_rate": "Cohort conversations with at least one human reply by the snapshot time divided by cohort conversations. Automated replies and bounces are excluded. Late replies can change a refreshed report.",
    "period_replies": "Human reply messages received during the selected dates, including replies to older tracked conversations.",
    "awaiting": "Cohort conversations with neither a human reply nor a recorded bounce by the snapshot time; no waiting-time threshold is implied.",
    "outcomes": "Interested, declined and won are operator-recorded. Unsubscribe is also detected from incoming opt-outs. Unclassified means no explicit outcome was recorded.",
    "reply_suggestions": "Automatic local-rule suggestions for the latest human reply in each cohort conversation. Suggestions are not confirmed outcomes or sales. Ambiguous wording stays needs_review. Saved snapshots retain their original suggestions.",
    "time": "Dates are inclusive in Africa/Johannesburg. Gmail internalDate dates replies; old records without it use local ingestion time.",
}


def instant(value):
    try:
        result = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return result.replace(tzinfo=timezone.utc) if result.tzinfo is None else result
    except (ValueError, TypeError):
        return None


def incoming_time(mail):
    try:
        return datetime.fromtimestamp(
            int(mail["provider_timestamp"]) / 1000, timezone.utc
        )
    except (ValueError, TypeError, KeyError, OverflowError, OSError):
        return instant(mail.get("created_at"))


def previous_week(at=None):
    today = (at or datetime.now(timezone.utc)).astimezone(TZ).date()
    monday = today - timedelta(days=today.weekday())
    return monday - timedelta(days=7), monday - timedelta(days=1)


def period(start=None, end=None):
    if start is None and end is None:
        start, end = previous_week()
    if start is None or end is None:
        raise ValueError("Provide both start and end dates")
    start = date.fromisoformat(str(start))
    end = date.fromisoformat(str(end))
    if end < start or (end - start).days > 365:
        raise ValueError("Choose an inclusive date range of 1 to 366 days")
    return start, end


def build_report(
    start=None, end=None, *, mailbox=None, include_tests=False, as_of=None
):
    from services.reply_service import assess_reply

    start, end = period(start, end)
    as_of = as_of or datetime.now(timezone.utc)
    left = datetime.combine(start, time.min, TZ)
    right = datetime.combine(end + timedelta(days=1), time.min, TZ)
    data = store.snapshot(("messages", "inbox", "outcomes", "leads", "gmail_cursor"))
    allowed = ("outreach", "test") if include_tests else ("outreach",)
    outbound = [
        m
        for m in data["messages"]
        if m.get("status") == "sent"
        and m.get("provider_message_id")
        and m.get("purpose") in allowed
        and (mailbox is None or m.get("from") == mailbox)
        and instant(m.get("created_at"))
        and instant(m["created_at"]) <= as_of
    ]
    outbound.sort(key=lambda m: (instant(m["created_at"]), m["id"]))
    groups = {}
    by_id = {}
    for m in outbound:
        key = (m.get("from"), m.get("thread_id") or m["id"])
        groups.setdefault(key, []).append(m)
        by_id[m["id"]] = key
    incoming = {}
    for mail in data["inbox"]:
        key = by_id.get(mail.get("matched_outbound_id"))
        if key is None:
            key = (mail.get("mailbox"), mail.get("thread_id"))
        # Require a tracked outbound, matching mailbox, and the expected participant.
        if key not in groups or mail.get("mailbox") != key[0]:
            continue
        if mail.get("sender") not in {
            m.get("to") for m in groups[key]
        } and not mail.get("bounce"):
            continue
        when = incoming_time(mail)
        if when is None or when > as_of or when < instant(groups[key][0]["created_at"]):
            continue
        identity = (mail.get("mailbox"), mail.get("provider_message_id") or mail["id"])
        incoming[identity] = (key, mail, when)
    outcomes = {}
    for event in sorted(data["outcomes"], key=lambda e: e.get("created_at", "")):
        when = instant(event.get("created_at"))
        if when and when <= as_of:
            outcomes[(event.get("mailbox"), event.get("thread_id"))] = event
    by_conversation = {}
    for key, mail, when in incoming.values():
        by_conversation.setdefault(key, []).append((mail, when))
    rows = []
    for key, messages in groups.items():
        first = messages[0]
        sent = instant(first["created_at"])
        if not left <= sent < right:
            continue
        replies = by_conversation.get(key, [])
        human = [
            (m, t) for m, t in replies if not m.get("automated") and not m.get("bounce")
        ]
        bounced = any(m.get("bounce") for m, _ in replies)
        event = outcomes.get(key, {})
        optouts = [(m, t) for m, t in human if m.get("unsubscribe")]
        outcome = event.get("outcome", "unclassified")
        outcome_note = event.get("note", "")
        # Classification must never hide a recorded opt-out.
        if optouts:
            outcome = "unsubscribed"
            latest_optout = max(optouts, key=lambda item: item[1])[0]
            outcome_note = "Recorded incoming opt-out: " + latest_optout["id"]
        latest_human = (
            max(human, key=lambda item: (item[1], item[0]["id"]))[0] if human else None
        )
        rows.append(
            {
                "id": first["id"],
                "mailbox": key[0],
                "thread_id": first.get("thread_id"),
                "recipient": first.get("to", ""),
                "company": first.get("company", ""),
                "subject": first.get("subject", ""),
                "sent_at": first["created_at"],
                "purpose": first["purpose"],
                "reply_count": len(human),
                "first_reply_at": (
                    min((t for _, t in human), default=None).isoformat()
                    if human
                    else None
                ),
                "bounced": bounced,
                "state": (
                    "responded" if human else "bounced" if bounced else "awaiting_reply"
                ),
                "outcome": outcome,
                "outcome_note": outcome_note,
                "reply_suggestion": (
                    assess_reply(latest_human)["category"] if latest_human else None
                ),
            }
        )
    cohort_replies = sum(bool(r["reply_count"]) for r in rows)
    period_human = [
        (m, t)
        for _, m, t in incoming.values()
        if left <= t < right and not m.get("automated") and not m.get("bounce")
    ]
    period_auto = [
        m
        for _, m, t in incoming.values()
        if left <= t < right and (m.get("automated") or m.get("bounce"))
    ]
    sends_in_period = [m for m in outbound if left <= instant(m["created_at"]) < right]
    counts = dict(Counter(r["outcome"] for r in rows))
    trend = []
    day = start
    while day <= end:
        trend.append(
            {
                "date": day.isoformat(),
                "sent": sum(
                    instant(m["created_at"]).astimezone(TZ).date() == day
                    for m in sends_in_period
                ),
                "human_replies": sum(
                    t.astimezone(TZ).date() == day for _, t in period_human
                ),
            }
        )
        day += timedelta(days=1)
    return {
        "start": start.isoformat(),
        "end": end.isoformat(),
        "timezone": str(TZ),
        "as_of": as_of.isoformat(),
        "mailbox": mailbox,
        "include_tests": include_tests,
        "summary": {
            "emails_sent": len(sends_in_period),
            "conversations_sent": len(rows),
            "conversations_replied": cohort_replies,
            "awaiting_reply": sum(r["state"] == "awaiting_reply" for r in rows),
            "bounced_conversations": sum(r["bounced"] for r in rows),
            "reply_rate_percent": (
                round(100 * cohort_replies / len(rows), 1) if rows else None
            ),
            "human_replies_received": len(period_human),
            "automated_messages_received": len(period_auto),
            "outcomes": {o: counts.get(o, 0) for o in OUTCOMES},
            "reply_suggestions": dict(
                Counter(r["reply_suggestion"] for r in rows if r["reply_suggestion"])
            ),
            "recorded_leads": len(data["leads"]),
        },
        "conversations": rows,
        "daily": trend,
        "definitions": DEFINITIONS,
        "last_sync_at": next(
            (
                c.get("updated_at")
                for c in data["gmail_cursor"]
                if c.get("id") == mailbox
            ),
            None,
        ),
        "reply_timestamp_fallbacks": sum(
            not m.get("provider_timestamp") for _, m, _ in incoming.values()
        ),
    }


def all_time_metrics():
    data = store.list("messages")
    dates = [
        instant(m.get("created_at")) for m in data if m.get("purpose") == "outreach"
    ]
    dates = [d for d in dates if d]
    # All-time spans are intentionally not restricted to the interactive export window.
    start = min(dates).astimezone(TZ).date() if dates else datetime.now(TZ).date()
    end = datetime.now(TZ).date()
    summaries = []
    while start <= end:
        stop = min(start + timedelta(days=365), end)
        summaries.append(build_report(start, stop)["summary"])
        start = stop + timedelta(days=1)
    sent = sum(s["emails_sent"] for s in summaries)
    total = sum(s["conversations_sent"] for s in summaries)
    replied = sum(s["conversations_replied"] for s in summaries)
    return {
        "totalEmailsSent": sent,
        "totalResponsesReceived": replied,
        "responseRatePercent": round(replied * 100 / total, 1) if total else None,
        "activeLeadsInPipeline": len(store.list("leads")),
        "replyTrackingAvailable": True,
    }


def safe_cell(value):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def workbook(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    for key in (
        "start",
        "end",
        "timezone",
        "as_of",
        "mailbox",
        "include_tests",
        "last_sync_at",
        "reply_timestamp_fallbacks",
    ):
        summary.append([key, safe_cell(report.get(key))])
    for key, value in report["summary"].items():
        if isinstance(value, dict):
            for outcome, count in value.items():
                prefix = "outcome" if key == "outcomes" else key
                summary.append([f"{prefix}_{outcome}", count])
        else:
            summary.append([key, value])
    sheet = wb.create_sheet("Conversations")
    columns = (
        "id",
        "mailbox",
        "recipient",
        "company",
        "subject",
        "sent_at",
        "purpose",
        "state",
        "reply_count",
        "first_reply_at",
        "bounced",
        "outcome",
        "outcome_note",
        "reply_suggestion",
    )
    sheet.append(list(columns))
    for row in report["conversations"]:
        sheet.append([safe_cell(row.get(c)) for c in columns])
    daily = wb.create_sheet("Daily")
    daily.append(["Date", "Confirmed sends", "Human reply messages"])
    for row in report["daily"]:
        daily.append([row["date"], row["sent"], row["human_replies"]])
    notes = wb.create_sheet("Definitions")
    for key, value in report["definitions"].items():
        notes.append([key, value])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                60, max(18, max(len(str(c.value or "")) for c in column) + 2)
            )
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def save_snapshot(start=None, end=None, *, mailbox, include_tests=False):
    report = build_report(start, end, mailbox=mailbox, include_tests=include_tests)
    return store.put("reports", uuid4().hex, report)


def report_attachment(report_id, mailbox):
    report = store.get("reports", report_id)
    if not report or report.get("mailbox") != mailbox:
        raise ValueError("Report snapshot does not belong to this mailbox")
    return {
        "filename": f"outreach-{report['start']}-to-{report['end']}.xlsx",
        "content": workbook(report),
    }


def queue_report(report):
    from agents.human_approval import queue_approval

    mailbox = report["mailbox"]
    s = report["summary"]
    return queue_approval(
        "report_" + report["id"],
        {
            "recipient_email": mailbox,
            "sender_mailbox": mailbox,
            "purpose": "report",
            "report_id": report["id"],
            "subject": f"Outreach report: {report['start']} to {report['end']}",
            "draft_response": f"Outreach report ({report['timezone']})\n\nPeriod: {report['start']} to {report['end']}\nConfirmed sends: {s['emails_sent']}\nConversations with a human reply: {s['conversations_replied']} / {s['conversations_sent']}\nAwaiting reply: {s['awaiting_reply']}\nHuman reply messages received in period: {s['human_replies_received']}\n\nThe attached workbook contains the saved snapshot and metric definitions. Gmail acceptance does not prove inbox delivery.\nSnapshot: {report['as_of']}",
        },
    )


async def generate_weekly_report():
    from services.email_service import email_service

    mailbox = email_service.expected_email
    if not mailbox:
        raise ValueError("Configure Gmail before generating a weekly report")
    start, end = previous_week()
    key = "weekly_" + sha256(f"{mailbox}:{start}:{end}".encode()).hexdigest()[:32]
    existing = store.get("reports", key)
    if existing:
        return existing
    report = build_report(start, end, mailbox=mailbox)
    return store.insert_once("reports", key, report)


async def send_weekly_report():
    # Compatibility name: create a reviewable draft, never silently send.
    from config import settings

    if not settings.WEEKLY_REPORTS_ENABLED:
        raise RuntimeError(
            "Weekly report scheduling is disabled; enable WEEKLY_REPORTS_ENABLED"
        )
    report = await generate_weekly_report()
    return {
        "report_id": report["id"],
        "approval": queue_report(report),
        "status": "awaiting_approval",
    }
