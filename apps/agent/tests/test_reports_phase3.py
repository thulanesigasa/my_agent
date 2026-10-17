"""Reporting invariants: real cohorts, time boundaries, stable exports and review-only delivery."""

from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from io import BytesIO
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from config import settings
from core.store import store
from services.report_service import (
    build_report,
    save_snapshot,
    workbook,
    previous_week,
    send_weekly_report,
    queue_report,
)
from services.email_service import email_service
from services.contacts import ensure_eligible, EligibilityError
from agents.human_approval import approve_draft
from test_foundation import gmail_client

MAILBOX = "sender@example.test"
AS_OF = datetime(2026, 10, 20, 10, tzinfo=timezone.utc)


def sent(record="one", when="2026-10-12T08:00:00+00:00", **values):
    result = store.put(
        "messages",
        record,
        {
            "status": "sent",
            "provider_message_id": "provider-" + record,
            "from": MAILBOX,
            "to": "owner@example.test",
            "thread_id": "thread-" + record,
            "purpose": "outreach",
            "subject": "Website",
            "body": "Hello",
            **values,
        },
    )
    # Seed original submission dates; StateStore.put intentionally preserves real timestamps.
    result["created_at"] = when
    with store.connect() as db:
        import json

        db.execute(
            "UPDATE records SET data=? WHERE kind='messages' AND id=?",
            (json.dumps(result), record),
        )
    return result


def reply(record="reply", outbound="one", when="2026-10-13T08:00:00+00:00", **values):
    return store.put(
        "inbox",
        record,
        {
            "mailbox": MAILBOX,
            "provider_message_id": record,
            "matched_outbound_id": outbound,
            "thread_id": "thread-" + outbound,
            "sender": "owner@example.test",
            "body": "Interested",
            "provider_timestamp": str(
                int(datetime.fromisoformat(when).timestamp() * 1000)
            ),
            "automated": False,
            "bounce": False,
            "status": "processed",
            **values,
        },
    )


def report(**kwargs):
    return build_report(
        "2026-10-12", "2026-10-18", mailbox=MAILBOX, as_of=AS_OF, **kwargs
    )


def test_empty_report_has_no_invented_counts_or_reply_rate():
    data = report()
    assert data["summary"]["emails_sent"] == 0
    assert data["summary"]["reply_rate_percent"] is None
    assert len(data["daily"]) == 7


def test_test_notifications_failed_unknown_and_other_mailboxes_excluded():
    sent()
    for record, values in (
        ("test", {"purpose": "test"}),
        ("notification", {"purpose": "notification"}),
        ("report", {"purpose": "report"}),
        ("failed", {"status": "failed"}),
        ("held", {"status": "delivery_unknown"}),
        ("other", {"from": "other@example.test"}),
        ("no-evidence", {"provider_message_id": None}),
    ):
        sent(record, **values)
    assert report()["summary"]["emails_sent"] == 1
    assert report(include_tests=True)["summary"]["emails_sent"] == 2


def test_replies_deduplicate_by_mailbox_provider_id_and_thread_is_one_cohort():
    sent()
    sent("another-send", thread_id="thread-one")
    reply()
    reply("duplicate", provider_message_id="reply")
    reply("second-reply")
    data = report()
    assert data["summary"]["emails_sent"] == 2
    assert data["summary"]["conversations_sent"] == 1
    assert data["summary"]["conversations_replied"] == 1
    assert data["summary"]["human_replies_received"] == 2
    assert data["summary"]["reply_rate_percent"] == 100


def test_late_replies_and_replies_to_older_sends_have_distinct_metrics():
    sent()
    sent("older", when="2026-10-05T08:00:00+00:00")
    reply(when="2026-10-19T08:00:00+00:00")
    reply("older-reply", outbound="older")
    data = report()
    assert data["summary"]["conversations_replied"] == 1
    assert data["summary"]["human_replies_received"] == 1
    assert data["summary"]["conversations_sent"] == 1


def test_automated_bounce_and_wrong_participant_never_count_as_human_reply():
    sent()
    reply(automated=True)
    reply("bounce", automated=True, bounce=True, sender="mailer-daemon@example.test")
    reply("attacker", sender="someone-else@example.test")
    data = report()
    assert data["summary"]["conversations_replied"] == 0
    assert data["summary"]["bounced_conversations"] == 1
    assert data["summary"]["awaiting_reply"] == 0
    assert data["summary"]["automated_messages_received"] == 2


def test_sast_boundaries_and_provider_timestamp_not_sync_time():
    sent("before", when="2026-10-11T21:59:59+00:00")
    sent("left", when="2026-10-11T22:00:00+00:00")
    sent("last", when="2026-10-18T21:59:59+00:00")
    sent("after", when="2026-10-18T22:00:00+00:00")
    reply(outbound="left", when="2026-10-11T22:01:00+00:00")
    data = report()
    assert data["summary"]["emails_sent"] == 2
    assert data["daily"][0] == {"date": "2026-10-12", "sent": 1, "human_replies": 1}
    assert data["daily"][-1]["sent"] == 1


def test_future_replies_do_not_change_snapshot_and_missing_dates_fallback():
    sent()
    reply(when="2026-10-21T08:00:00+00:00")
    assert report()["summary"]["conversations_replied"] == 0
    reply("fallback", provider_timestamp=None, created_at="2026-10-14T10:00:00+00:00")
    # put sets ingestion timestamp, which is correctly used for a legacy reply.
    assert report()["reply_timestamp_fallbacks"] == 1


def test_outcome_is_explicit_and_optout_overrides_older_classification():
    sent()
    reply()
    assert report()["summary"]["outcomes"]["unclassified"] == 1
    store.record_outcome(
        MAILBOX, "thread-one", "owner@example.test", "interested", "Asked for pricing"
    )
    assert report()["summary"]["outcomes"]["interested"] == 1
    reply("optout", when="2026-10-19T12:00:00+00:00", unsubscribe=True)
    assert report()["summary"]["outcomes"]["unsubscribed"] == 1


def test_snapshots_are_unchanged_by_later_replies_and_exports_escape_formulas():
    sent(subject='=HYPERLINK("bad")', company="  +SUM(1,1)")
    saved = save_snapshot("2026-10-12", "2026-10-18", mailbox=MAILBOX)
    reply()
    assert store.get("reports", saved["id"])["summary"]["conversations_replied"] == 0
    assert report()["summary"]["conversations_replied"] == 1
    wb = load_workbook(BytesIO(workbook(saved)))
    assert wb.sheetnames == ["Summary", "Conversations", "Daily", "Definitions"]
    assert wb["Conversations"]["E2"].data_type == "s"
    assert wb["Conversations"]["E2"].value.startswith("'=")
    assert wb["Conversations"]["D2"].value.startswith("'  +")


def test_period_validation_and_previous_complete_week():
    assert tuple(
        map(str, previous_week(datetime(2026, 10, 19, 5, tzinfo=timezone.utc)))
    ) == ("2026-10-12", "2026-10-18")
    assert tuple(
        map(str, previous_week(datetime(2026, 10, 18, 22, tzinfo=timezone.utc)))
    ) == ("2026-10-12", "2026-10-18")
    for start, end in (
        ("2026-10-18", None),
        ("2026-10-18", "2026-10-12"),
        ("2020-01-01", "2026-10-12"),
    ):
        with pytest.raises(ValueError):
            build_report(start, end)


@pytest.fixture
def api():
    from main import app

    with TestClient(app) as client:
        yield client


def test_report_routes_auth_validation_mailbox_isolation_and_binary_download(
    api, monkeypatch
):
    sent()
    monkeypatch.setattr(settings, "API_KEY_REQUIRED", True)
    assert api.get("/api/reports/summary").status_code == 401
    monkeypatch.setattr(settings, "API_KEY_REQUIRED", False)
    assert api.get("/api/reports/summary?start=bad").status_code == 422
    assert (
        api.get("/api/reports/summary?start=2026-10-18&end=2026-10-12").status_code
        == 400
    )
    saved = api.post(
        "/api/reports/snapshots", json={"start": "2026-10-12", "end": "2026-10-18"}
    ).json()
    response = api.get(f"/api/reports/snapshots/{saved['id']}/export")
    assert response.status_code == 200
    assert response.content[:2] == b"PK"
    monkeypatch.setattr(email_service, "expected_email", "different@example.test")
    assert api.get(f"/api/reports/snapshots/{saved['id']}/export").status_code == 404
    assert api.get("/api/reports/snapshots").json()["reports"] == []


def test_outcome_requires_human_evidence_and_unsubscribe_blocks_contact(api):
    sent()
    path = "/api/reports/conversations/one/outcome"
    assert (
        api.put(path, json={"outcome": "interested", "note": "No reply"}).status_code
        == 400
    )
    reply()
    assert (
        api.put(
            path, json={"outcome": "interested", "note": "Asked for pricing"}
        ).status_code
        == 200
    )
    assert api.put(path, json={"outcome": "interested", "note": " "}).status_code == 400
    assert (
        api.put(
            path, json={"outcome": "unsubscribed", "note": "Asked us to stop"}
        ).status_code
        == 200
    )
    with pytest.raises(EligibilityError):
        ensure_eligible("owner@example.test", "outreach")


@pytest.mark.asyncio
async def test_weekly_job_creates_stable_snapshot_and_one_reviewable_draft(monkeypatch):
    monkeypatch.setattr(settings, "WEEKLY_REPORTS_ENABLED", True)
    send = Mock(side_effect=AssertionError("No automatic email"))
    monkeypatch.setattr(email_service, "send_email", send)
    first = await send_weekly_report()
    second = await send_weekly_report()
    assert first["report_id"] == second["report_id"]
    assert len(store.list("reports")) == len(store.list("approvals")) == 1
    assert first["approval"]["recipient_email"] == MAILBOX
    send.assert_not_called()


def test_atomic_snapshot_creation_across_workers():
    with ThreadPoolExecutor(max_workers=4) as workers:
        reports = list(
            workers.map(
                lambda _: store.insert_once("reports", "one-week", {"value": "same"}),
                range(4),
            )
        )
    assert len(store.list("reports")) == 1
    assert len({r["created_at"] for r in reports}) == 1


@pytest.mark.asyncio
async def test_approved_report_sends_excel_attachment_once_and_never_counts_as_outreach(
    monkeypatch,
):
    saved = save_snapshot("2026-10-12", "2026-10-18", mailbox=MAILBOX)
    approval = queue_report(saved)
    client, messages = gmail_client()
    monkeypatch.setattr(email_service, "client", client)
    await approve_draft(approval["id"])
    await approve_draft(approval["id"])
    assert messages.send.call_count == 1
    import base64
    from email import message_from_bytes

    msg = message_from_bytes(
        base64.urlsafe_b64decode(messages.send.call_args.kwargs["body"]["raw"])
    )
    assert msg["To"] == MAILBOX
    parts = list(msg.walk())
    attachment = next(p for p in parts if p.get_filename())
    assert attachment.get_filename().endswith(".xlsx")
    wb = load_workbook(BytesIO(attachment.get_payload(decode=True)))
    assert "Conversations" in wb.sheetnames
    assert store.metrics()["totalEmailsSent"] == 0


@pytest.mark.asyncio
async def test_report_cannot_target_unrelated_mailbox_or_attach_other_snapshot():
    from agents.human_approval import queue_approval

    saved = save_snapshot("2026-10-12", "2026-10-18", mailbox="other@example.test")
    item = queue_approval(
        "wrong",
        {
            "recipient_email": MAILBOX,
            "subject": "Report",
            "draft_response": "See attached",
            "purpose": "report",
            "report_id": saved["id"],
        },
    )
    with pytest.raises(ValueError):
        await approve_draft(item["id"])
    with pytest.raises(EligibilityError):
        ensure_eligible("not-me@example.test", "report")
    assert store.get("approvals", "wrong")["status"] == "pending"


def test_weekly_schedule_is_opt_in_and_uses_monday_sast(monkeypatch):
    from scheduler import AgentScheduler

    monkeypatch.setattr(settings, "WEEKLY_REPORTS_ENABLED", True)
    scheduler = AgentScheduler()
    scheduler.scheduler = Mock()
    scheduler.start()
    calls = scheduler.scheduler.add_job.call_args_list
    weekly = next(c for c in calls if c.kwargs.get("id") == "weekly_report")
    assert str(weekly.kwargs["trigger"].timezone) == "Africa/Johannesburg"
    assert "mon" in str(weekly.kwargs["trigger"])
    assert "hour='7'" in str(weekly.kwargs["trigger"])
    monkeypatch.setattr(settings, "WEEKLY_REPORTS_ENABLED", False)
    scheduler.scheduler.reset_mock()
    scheduler.start()
    assert not any(
        c.kwargs.get("id") == "weekly_report"
        for c in scheduler.scheduler.add_job.call_args_list
    )


@pytest.mark.asyncio
async def test_held_report_send_never_resubmits_attachment(monkeypatch):
    from core.store import ConflictError
    from services.email_service import DeliveryUnknown

    saved = save_snapshot("2026-10-12", "2026-10-18", mailbox=MAILBOX)
    approval = queue_report(saved)
    client, messages = gmail_client(error=TimeoutError("Unknown submission outcome"))
    monkeypatch.setattr(email_service, "client", client)
    with pytest.raises(DeliveryUnknown):
        await approve_draft(approval["id"])
    with pytest.raises(ConflictError):
        await approve_draft(approval["id"])
    assert messages.send.call_count == 1
    assert store.get("approvals", approval["id"])["status"] == "delivery_unknown"
