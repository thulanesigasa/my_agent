"""Approval-only coordinator boundaries using isolated state and fake sync."""

import asyncio
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from unittest.mock import AsyncMock

import pytest
from fastapi.testclient import TestClient
from config import settings
from core.store import store, ConflictError
from services import automation as auto, research_jobs
from services.contacts import suppress
from services.email_service import email_service
from services.lead_discovery import DiscoveryUnavailable
from services.lead_service import ingest_candidate
from services.report_service import TZ, build_report, workbook
from test_leads_phase4 import qualified, candidate
from test_followups_phase5 import original


def enable(**changes):
    payload = {
        **auto.DEFAULT,
        "version": auto.config(email_service.expected_email)["version"],
        "paused": False,
        "local_time": "00:00",
        "weekdays_only": False,
        "weekly_report_email": False,
        **changes,
    }
    return auto.update(auto.AutomationInput(**payload), email_service.expected_email)


def allow_next_tick():
    store.delete("automation_runtime", email_service.expected_email)


def incoming(body="I'm interested. Please send more details.", **changes):
    return store.put(
        "inbox",
        "incoming",
        {
            "mailbox": email_service.expected_email,
            "sender": "owner@example.test",
            "thread_id": "thread-1",
            "rfc_message_id": "<incoming@example.test>",
            "matched_outbound_id": "initial",
            "subject": "Re: Website proposal",
            "body": body,
            "provider_timestamp": str(
                int(datetime.now(timezone.utc).timestamp() * 1000)
            ),
            "status": "processed",
            "automated": False,
            "bounce": False,
            "unsubscribe": False,
            **changes,
        },
    )


@pytest.fixture(autouse=True)
def providers(monkeypatch):
    async def sync():
        store.put(
            "gmail_cursor",
            email_service.expected_email,
            {"thread_ids": [m.get("thread_id") for m in store.list("messages")]},
        )
        return {"recorded": 0}

    sync_mock = AsyncMock(side_effect=sync)
    monkeypatch.setattr(auto, "process_inbox_batch", sync_mock)
    monkeypatch.setattr(
        email_service,
        "send_email",
        AsyncMock(side_effect=AssertionError("Must never send automatically")),
    )
    return sync_mock


@pytest.mark.asyncio
async def test_paused_default_does_no_work_or_provider_calls(providers):
    assert auto.listing(email_service.expected_email)["config"]["paused"]
    assert await auto.tick() is None
    providers.assert_not_awaited()
    assert store.list("approvals") == store.list("reports") == []


@pytest.mark.asyncio
async def test_eligible_initial_and_weekly_report_are_queued_not_sent(providers):
    lead = qualified()
    enable(weekly_report_email=True)
    run = await auto.tick()
    assert run["status"] == "completed"
    approvals = store.list("approvals")
    assert {a["purpose"] for a in approvals} == {"outreach", "report"}
    assert all(a["status"] == "pending" for a in approvals)
    initial = next(a for a in approvals if a["purpose"] == "outreach")
    assert initial["lead_id"] == lead["id"]
    assert (
        initial["proposal_snapshot"]["portfolio_url"] == "https://ts-industries.co.za/"
    )
    assert store.list("messages") == []
    assert auto.usage(email_service.expected_email)["used"] == 1
    assert await auto.tick() is None
    providers.assert_awaited_once()
    allow_next_tick()
    await auto.tick()
    assert len(store.list("approvals")) == 2
    assert len(store.list("reports")) == 1
    assert auto.usage(email_service.expected_email)["used"] == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("block", ["unreviewed", "permission", "suppressed"])
async def test_missing_evidence_or_permission_is_skipped_without_inventing_it(block):
    if block == "unreviewed":
        lead = ingest_candidate(candidate())
    else:
        lead = qualified(consent=block != "permission")
        if block == "suppressed":
            suppress("owner@example.test", "Do not contact")
    before = store.get("leads", lead["id"])
    contact = store.get("contacts", "owner@example.test")
    enable()
    run = await auto.tick()
    assert run["lead_exceptions"][0]["lead_id"] == lead["id"]
    assert store.get("leads", lead["id"]) == before
    assert store.get("contacts", "owner@example.test") == contact
    assert store.list("approvals") == []


@pytest.mark.asyncio
async def test_supported_latest_business_reply_is_drafted_and_test_mail_is_excluded():
    original()
    incoming()
    enable()
    await auto.tick()
    replies = [a for a in store.list("approvals") if a.get("incoming_id")]
    assert len(replies) == 1 and replies[0]["provider_thread_id"] == "thread-1"
    assert replies[0]["status"] == "pending"
    assert store.list("outcomes") == []
    assert len(store.list("automation_reply_assessments")) == 1
    allow_next_tick()
    await auto.tick()
    assert len([a for a in store.list("approvals") if a.get("incoming_id")]) == 1


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "changes",
    [
        {"automated": True},
        {"bounce": True},
        {"unsubscribe": True},
        {"mailbox": "other@example.test"},
        {"conversation_purpose": "test"},
    ],
)
async def test_nonbusiness_or_unsafe_incoming_does_not_get_reply(changes):
    original()
    incoming(**changes)
    enable()
    await auto.tick()
    assert not any(a.get("incoming_id") for a in store.list("approvals"))


@pytest.mark.asyncio
async def test_followups_run_after_automatic_sync_without_enabling_separate_schedule():
    original()
    enable()
    await auto.tick()
    followups = [a for a in store.list("approvals") if a.get("followup_of")]
    assert len(followups) == 1 and followups[0]["status"] == "pending"
    assert store.list("followup_schedules") == []
    assert auto.usage(email_service.expected_email)["used"] == 1


@pytest.mark.asyncio
async def test_failed_sync_blocks_replies_and_followups_but_saves_report(providers):
    original()
    providers.side_effect = RuntimeError("Provider outage")
    enable()
    result = await auto.tick()
    assert result["status"] == "completed"
    assert result["results"][0]["status"] == "blocked"
    assert len(store.list("approvals")) == 1  # Original sent approval only.
    assert len(store.list("reports")) == 1


def test_shared_daily_slot_budget_is_atomic_and_cannot_reset_by_edit():
    saved = enable(daily_draft_limit=3)
    mailbox = email_service.expected_email
    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(
            pool.map(
                lambda i: auto._reserve(mailbox, "initial", str(i), saved["version"]),
                range(12),
            )
        )
    assert sum(results) == 3
    enable(daily_draft_limit=1)
    assert auto.usage(mailbox)["remaining"] == 0
    latest = enable(daily_draft_limit=3)
    assert not auto._reserve(mailbox, "reply", "new", latest["version"])
    assert auto.usage("other@example.test")["used"] == 0


@pytest.mark.asyncio
async def test_concurrent_ticks_and_poll_interval_do_not_repeat_sync(providers):
    enable()
    results = await asyncio.gather(*(auto.tick() for _ in range(5)))
    assert sum(r is not None for r in results) == 1
    providers.assert_awaited_once()


@pytest.mark.asyncio
async def test_pause_during_sync_stops_later_work_and_pauses_owned_research(providers):
    qualified()
    enable()
    job_id = auto.research_prefix(email_service.expected_email) + "today"
    store.put("research_jobs", job_id, {"status": "queued"})

    async def pause():
        enable(paused=True)
        return {}

    providers.side_effect = pause
    result = await auto.tick()
    assert result["status"] == "stopped"
    assert store.get("research_jobs", job_id)["status"] == "paused"
    assert store.list("approvals") == []


@pytest.mark.asyncio
async def test_crash_after_queue_pauses_and_recovery_never_overwrites_approval(
    monkeypatch,
):
    qualified()
    enable()
    draft = auto.draft_lead

    def crash(*args):
        draft(*args)
        raise RuntimeError("Crash after persisted approval")

    monkeypatch.setattr(auto, "draft_lead", crash)
    with pytest.raises(RuntimeError):
        await auto.tick()
    assert auto.config(email_service.expected_email)["paused"]
    approval = store.list("approvals")[0]
    run = store.list("automation_runs")[0]
    store.put("automation_runs", run["id"], {**run, "status": "running"})
    enable()
    allow_next_tick()
    assert await auto.tick() is None
    assert auto.config(email_service.expected_email)["paused"]
    assert store.get("automation_runs", run["id"])["status"] == "interrupted"
    assert store.list("approvals") == [approval]


@pytest.mark.asyncio
async def test_research_is_one_deterministic_job_daily_and_provider_optin_is_required(
    monkeypatch,
):
    monkeypatch.setattr(settings, "RESEARCH_JOBS_ENABLED", True)
    monkeypatch.setattr(settings, "ALLOW_TAVILY_RESEARCH", True)
    monkeypatch.setattr(settings, "TAVILY_API_KEY", "fixture-only")
    enable(research_enabled=True, location="Example Town", industry="plumbers")
    await auto.tick()
    allow_next_tick()
    await auto.tick()
    jobs = store.list("research_jobs")
    assert len(jobs) == 1 and jobs[0]["max_requests"] == 3
    # A deterministic ID remains idempotent even after a response was lost.
    assert (
        research_jobs.create_job("Example Town", "plumbers", job_id=jobs[0]["id"])["id"]
        == jobs[0]["id"]
    )
    with pytest.raises(ConflictError):
        research_jobs.create_job("Different Town", "plumbers", job_id=jobs[0]["id"])


@pytest.mark.asyncio
async def test_missing_research_configuration_is_visible_without_stopping_sync_and_reports(
    monkeypatch,
):
    enable(research_enabled=True, location="Example Town", industry="plumbers")

    def unavailable(*args, **kwargs):
        raise DiscoveryUnavailable("Research provider not configured")

    monkeypatch.setattr(research_jobs, "create_job", unavailable)
    result = await auto.tick()
    assert result["status"] == "completed"
    assert any(
        r["stage"] == "research" and r["status"] == "blocked" for r in result["results"]
    )
    assert not auto.config(email_service.expected_email)["paused"]


def test_report_suggestions_do_not_change_confirmed_outcomes_and_export_separately():
    original()
    incoming("I'm interested")
    start = (datetime.now(TZ) - timedelta(days=9)).date()
    end = datetime.now(TZ).date()
    report = build_report(start, end, mailbox=email_service.expected_email)
    assert report["summary"]["reply_suggestions"] == {"interested": 1}
    assert report["summary"]["outcomes"]["interested"] == 0
    assert report["summary"]["outcomes"]["unclassified"] == 1
    incoming("I'm interested but not now")
    latest = build_report(start, end, mailbox=email_service.expected_email)
    assert latest["summary"]["reply_suggestions"] == {"needs_review": 1}
    from io import BytesIO
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(workbook(report)))
    rows = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True)}
    assert rows["reply_suggestions_interested"] == 1
    assert rows["outcome_interested"] == 0


def test_api_auth_validation_version_and_mailbox_scope(monkeypatch):
    from main import app

    with TestClient(app) as api:
        monkeypatch.setattr(settings, "API_KEY_REQUIRED", True)
        assert api.get("/api/automation").status_code == 401
        assert api.patch("/api/automation", json={}).status_code == 401
        monkeypatch.setattr(settings, "API_KEY_REQUIRED", False)
        payload = {**auto.DEFAULT, "paused": False}
        assert (
            api.patch(
                "/api/automation", json={**payload, "send_automatically": True}
            ).status_code
            == 422
        )
        assert (
            api.patch(
                "/api/automation", json={**payload, "research_enabled": True}
            ).status_code
            == 422
        )
        assert api.patch("/api/automation", json=payload).status_code == 200
        assert api.patch("/api/automation", json=payload).status_code == 409
        monkeypatch.setattr(email_service, "expected_email", "other@example.test")
        data = api.get("/api/automation").json()
        assert data["config"]["paused"] and data["runs"] == []


@pytest.mark.asyncio
async def test_runner_shutdown_waits_for_tick(monkeypatch):
    entered = asyncio.Event()

    async def tick(stop):
        entered.set()

    monkeypatch.setattr(auto, "tick", tick)
    runner = auto.AutomationRunner()
    runner.start()
    await asyncio.wait_for(entered.wait(), timeout=2)
    await runner.stop()
    assert runner.task is None


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "changes", [{"location": "Different Town"}, {"industry": "plumbers"}]
)
async def test_target_filters_constrain_initial_outreach(changes):
    qualified()
    enable(**changes)
    await auto.tick()
    assert store.list("approvals") == []


@pytest.mark.asyncio
async def test_stop_flag_prevents_even_sync(providers):
    enable()
    assert await auto.tick(lambda: True) is None
    providers.assert_not_awaited()


def test_window_uses_sast_weekday_and_time():
    saved = {**auto.DEFAULT, "local_time": "09:00"}
    assert not auto._window(saved, datetime(2026, 10, 19, 6, 59, tzinfo=timezone.utc))
    assert auto._window(saved, datetime(2026, 10, 19, 7, 0, tzinfo=timezone.utc))
    assert not auto._window(saved, datetime(2026, 10, 18, 10, 0, tzinfo=timezone.utc))
